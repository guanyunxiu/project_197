from rest_framework import viewsets, status, permissions, filters
from rest_framework.decorators import action, api_view, permission_classes, parser_classes
from rest_framework.response import Response
from rest_framework.permissions import AllowAny, IsAuthenticated, IsAdminUser
from rest_framework_simplejwt.tokens import RefreshToken
from rest_framework.parsers import MultiPartParser, FormParser
from django.contrib.auth import authenticate
from django.db.models import Count, Sum, Q, F, ExpressionWrapper, IntegerField
from django.db.models.functions import TruncMonth, TruncDate
from django.shortcuts import get_object_or_404
from django.utils import timezone
from datetime import timedelta
from django.conf import settings
import openpyxl
from io import BytesIO

from .models import Category, Book, BorrowRecord, Reservation, RenewRecord, FineRecord, OperationLog, SystemConfig
from .serializers import (
    UserSerializer, UserRegistrationSerializer, UserLoginSerializer,
    CategorySerializer, BookSerializer, BookListSerializer,
    BorrowRecordSerializer, BorrowBookSerializer, ReturnBookSerializer,
    ChangePasswordSerializer, ReservationSerializer, CreateReservationSerializer,
    RenewRecordSerializer, CreateRenewSerializer, ReviewRenewSerializer,
    FineRecordSerializer, PayFineSerializer, OperationLogSerializer,
    SystemConfigSerializer
)
from .utils import log_operation, get_system_config
from django.contrib.auth import get_user_model
from decimal import Decimal

User = get_user_model()


class IsSuperAdmin(permissions.BasePermission):
    def has_permission(self, request, view):
        return request.user.is_authenticated and request.user.role == 'super_admin'


class IsAdminRole(permissions.BasePermission):
    def has_permission(self, request, view):
        return request.user.is_authenticated and request.user.role in ['super_admin', 'admin']


class IsAdminOrSelf(permissions.BasePermission):
    def has_object_permission(self, request, view, obj):
        if request.user.role in ['super_admin', 'admin']:
            return True
        return obj == request.user


class IsAdminOrOwner(permissions.BasePermission):
    def has_object_permission(self, request, view, obj):
        if request.user.role in ['super_admin', 'admin']:
            return True
        return getattr(obj, 'reader', None) == request.user


@api_view(['POST'])
@permission_classes([AllowAny])
def register(request):
    serializer = UserRegistrationSerializer(data=request.data)
    if serializer.is_valid():
        user = serializer.save()
        user.role = 'reader'
        user.save()
        refresh = RefreshToken.for_user(user)

        log_operation(
            request, 'create', 'User', user.id, user.username,
            f'用户注册: {user.username}'
        )

        return Response({
            'user': UserSerializer(user).data,
            'refresh': str(refresh),
            'access': str(refresh.access_token),
        }, status=status.HTTP_201_CREATED)
    return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)


@api_view(['POST'])
@permission_classes([AllowAny])
def login(request):
    serializer = UserLoginSerializer(data=request.data)
    if serializer.is_valid():
        user = authenticate(
            username=serializer.validated_data['username'],
            password=serializer.validated_data['password']
        )
        if user:
            if not user.is_active:
                return Response({'detail': '账号已被禁用'}, status=status.HTTP_403_FORBIDDEN)

            refresh = RefreshToken.for_user(user)

            log_operation(
                request, 'login', 'User', user.id, user.username,
                f'用户登录: {user.username}'
            )

            return Response({
                'user': UserSerializer(user).data,
                'refresh': str(refresh),
                'access': str(refresh.access_token),
            })
        return Response({'detail': '用户名或密码错误'}, status=status.HTTP_401_UNAUTHORIZED)
    return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)


@api_view(['POST'])
@permission_classes([IsAuthenticated])
def logout(request):
    try:
        refresh_token = request.data.get('refresh')
        if refresh_token:
            token = RefreshToken(refresh_token)
            token.blacklist()

        log_operation(
            request, 'logout', 'User', request.user.id, request.user.username,
            f'用户退出: {request.user.username}'
        )

        return Response({'detail': '退出成功'}, status=status.HTTP_200_OK)
    except Exception:
        return Response({'detail': '退出成功'}, status=status.HTTP_200_OK)


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def get_current_user(request):
    return Response(UserSerializer(request.user).data)


@api_view(['POST'])
@permission_classes([IsAuthenticated])
def change_password(request):
    serializer = ChangePasswordSerializer(data=request.data)
    if serializer.is_valid():
        if not request.user.check_password(serializer.validated_data['old_password']):
            return Response({'old_password': '旧密码错误'}, status=status.HTTP_400_BAD_REQUEST)
        request.user.set_password(serializer.validated_data['new_password'])
        request.user.save()

        log_operation(
            request, 'update', 'User', request.user.id, request.user.username,
            '修改密码'
        )

        return Response({'detail': '密码修改成功'}, status=status.HTTP_200_OK)
    return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)


class UserViewSet(viewsets.ModelViewSet):
    queryset = User.objects.all()
    serializer_class = UserSerializer
    permission_classes = [IsAuthenticated]
    filter_backends = [filters.SearchFilter, filters.OrderingFilter]
    search_fields = ['username', 'first_name', 'last_name', 'email', 'phone']
    ordering_fields = ['date_joined', 'created_at', 'username']

    def get_permissions(self):
        if self.action in ['create']:
            return [AllowAny()]
        elif self.action in ['list', 'destroy']:
            return [IsAdminRole()]
        return [IsAuthenticated(), IsAdminOrSelf()]

    def get_queryset(self):
        queryset = User.objects.all()
        role = self.request.query_params.get('role')
        if role:
            queryset = queryset.filter(role=role)
        search = self.request.query_params.get('search')
        if search:
            queryset = queryset.filter(
                Q(username__icontains=search) |
                Q(first_name__icontains=search) |
                Q(email__icontains=search)
            )
        return queryset

    def update(self, request, *args, **kwargs):
        partial = kwargs.pop('partial', False)
        instance = self.get_object()

        if request.user.role not in ['super_admin', 'admin'] and instance != request.user:
            return Response({'detail': '无权限修改'}, status=status.HTTP_403_FORBIDDEN)

        if request.user.role != 'super_admin' and instance.role == 'super_admin':
            return Response({'detail': '无权限修改超级管理员'}, status=status.HTTP_403_FORBIDDEN)

        data = request.data.copy()
        if request.user.role not in ['super_admin', 'admin']:
            data.pop('role', None)
            data.pop('is_active', None)
        if request.user.role == 'admin' and 'role' in data and data['role'] == 'super_admin':
            return Response({'detail': '无权限提升为超级管理员'}, status=status.HTTP_403_FORBIDDEN)

        serializer = self.get_serializer(instance, data=data, partial=partial)
        serializer.is_valid(raise_exception=True)
        self.perform_update(serializer)

        log_operation(
            request, 'update', 'User', instance.id, instance.username,
            f'更新用户信息: {instance.username}'
        )

        return Response(serializer.data)

    def destroy(self, request, *args, **kwargs):
        instance = self.get_object()
        if instance.role == 'super_admin' and request.user.role != 'super_admin':
            return Response({'detail': '无权限删除超级管理员'}, status=status.HTTP_403_FORBIDDEN)

        log_operation(
            request, 'delete', 'User', instance.id, instance.username,
            f'删除用户: {instance.username}'
        )

        return super().destroy(request, *args, **kwargs)


class CategoryViewSet(viewsets.ModelViewSet):
    queryset = Category.objects.annotate(book_count=Count('books'))
    serializer_class = CategorySerializer
    permission_classes = [IsAuthenticated]
    filter_backends = [filters.SearchFilter, filters.OrderingFilter]
    search_fields = ['name', 'description']
    ordering_fields = ['name', 'created_at', 'book_count']

    def get_permissions(self):
        if self.action in ['list', 'retrieve']:
            return [IsAuthenticated()]
        return [IsAdminRole()]

    def get_queryset(self):
        queryset = Category.objects.annotate(book_count=Count('books'))
        search = self.request.query_params.get('search')
        if search:
            queryset = queryset.filter(name__icontains=search)
        return queryset

    def perform_create(self, serializer):
        instance = serializer.save()
        log_operation(
            self.request, 'create', 'Category', instance.id, instance.name,
            f'创建分类: {instance.name}'
        )

    def perform_update(self, serializer):
        instance = serializer.save()
        log_operation(
            self.request, 'update', 'Category', instance.id, instance.name,
            f'更新分类: {instance.name}'
        )

    def perform_destroy(self, instance):
        log_operation(
            self.request, 'delete', 'Category', instance.id, instance.name,
            f'删除分类: {instance.name}'
        )
        instance.delete()


class BookViewSet(viewsets.ModelViewSet):
    queryset = Book.objects.all()
    serializer_class = BookSerializer
    permission_classes = [IsAuthenticated]
    filter_backends = [filters.SearchFilter, filters.OrderingFilter]
    search_fields = ['title', 'author', 'isbn', 'publisher']
    ordering_fields = ['title', 'author', 'borrow_count', 'created_at', 'available_quantity']

    def get_permissions(self):
        if self.action in ['list', 'retrieve', 'reserve', 'hot_books']:
            return [IsAuthenticated()]
        return [IsAdminRole()]

    def get_serializer_class(self):
        if self.action == 'list':
            return BookListSerializer
        return BookSerializer

    def get_serializer_context(self):
        context = super().get_serializer_context()
        context['request'] = self.request
        return context

    def get_queryset(self):
        queryset = Book.objects.all()

        category = self.request.query_params.get('category')
        if category:
            queryset = queryset.filter(category_id=category)

        book_status = self.request.query_params.get('status')
        if book_status:
            queryset = queryset.filter(status=book_status)

        search_type = self.request.query_params.get('search_type')
        search = self.request.query_params.get('search')
        if search:
            if search_type == 'title':
                queryset = queryset.filter(title__icontains=search)
            elif search_type == 'author':
                queryset = queryset.filter(author__icontains=search)
            elif search_type == 'isbn':
                queryset = queryset.filter(isbn__exact=search)
            else:
                queryset = queryset.filter(
                    Q(title__icontains=search) |
                    Q(author__icontains=search) |
                    Q(isbn__icontains=search) |
                    Q(publisher__icontains=search)
                )

        ordering = self.request.query_params.get('ordering', '-borrow_count')
        if ordering == 'hot':
            queryset = queryset.order_by('-borrow_count', '-created_at')
        elif ordering == 'newest':
            queryset = queryset.order_by('-created_at')
        elif ordering == 'name':
            queryset = queryset.order_by('title')
        else:
            queryset = queryset.order_by(ordering)

        return queryset

    @action(detail=False, methods=['get'])
    def hot_books(self, request):
        books = Book.objects.all().order_by('-borrow_count')[:10]
        serializer = BookListSerializer(books, many=True, context={'request': request})
        return Response(serializer.data)

    @action(detail=True, methods=['post'], permission_classes=[IsAuthenticated])
    def reserve(self, request, pk=None):
        book = self.get_object()
        serializer = CreateReservationSerializer(data={'book_id': book.id}, context={'request': request})
        if serializer.is_valid():
            reservation = serializer.save()

            log_operation(
                request, 'reserve', 'Reservation', reservation.id, book.title,
                f'预约图书: {book.title}'
            )

            return Response(ReservationSerializer(reservation).data, status=status.HTTP_201_CREATED)
        return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)

    @action(detail=False, methods=['post'], permission_classes=[IsAdminRole])
    def upload_cover(self, request):
        cover_file = request.FILES.get('file')
        if not cover_file:
            return Response({'detail': '请上传封面图片'}, status=status.HTTP_400_BAD_REQUEST)
        
        try:
            from PIL import Image
            img = Image.open(cover_file)
            img.verify()
        except Exception:
            return Response({'detail': '无效的图片文件'}, status=status.HTTP_400_BAD_REQUEST)
        
        if cover_file.size > 2 * 1024 * 1024:
            return Response({'detail': '图片大小不能超过2MB'}, status=status.HTTP_400_BAD_REQUEST)
        
        import os
        from django.conf import settings
        from django.core.files.storage import default_storage
        from django.core.files.base import ContentFile
        
        ext = os.path.splitext(cover_file.name)[1].lower()
        if ext not in ['.jpg', '.jpeg', '.png', '.gif', '.webp']:
            return Response({'detail': '不支持的图片格式'}, status=status.HTTP_400_BAD_REQUEST)
        
        import uuid
        filename = f'book_covers/{uuid.uuid4()}{ext}'
        saved_path = default_storage.save(filename, ContentFile(cover_file.read()))
        cover_url = default_storage.url(saved_path)
        
        log_operation(
            request, 'update', 'Book', None, cover_file.name,
            f'上传图书封面: {cover_file.name}'
        )
        
        return Response({
            'url': cover_url,
            'path': saved_path
        }, status=status.HTTP_201_CREATED)

    @action(detail=False, methods=['post'], permission_classes=[IsAdminRole])
    def import_excel(self, request):
        excel_file = request.FILES.get('file')
        if not excel_file:
            return Response({'detail': '请上传Excel文件'}, status=status.HTTP_400_BAD_REQUEST)

        try:
            wb = openpyxl.load_workbook(BytesIO(excel_file.read()))
            ws = wb.active

            headers = [cell.value for cell in ws[1]]
            required_fields = ['isbn', 'title', 'author', 'publisher', 'publish_date', 'category', 'description', 'location', 'total_quantity']

            missing_fields = [f for f in required_fields if f not in headers]
            if missing_fields:
                return Response(
                    {'detail': f'缺少必要列: {", ".join(missing_fields)}'},
                    status=status.HTTP_400_BAD_REQUEST
                )

            created_count = 0
            updated_count = 0
            errors = []

            for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                row_data = dict(zip(headers, row))

                try:
                    category_name = row_data.get('category')
                    category, _ = Category.objects.get_or_create(name=category_name)

                    book_data = {
                        'isbn': str(row_data.get('isbn', '')),
                        'title': str(row_data.get('title', '')),
                        'author': str(row_data.get('author', '')),
                        'publisher': str(row_data.get('publisher', '')),
                        'publish_date': row_data.get('publish_date'),
                        'category': category.id,
                        'description': str(row_data.get('description', '')),
                        'location': str(row_data.get('location', '')),
                        'total_quantity': int(row_data.get('total_quantity', 1)),
                        'available_quantity': int(row_data.get('available_quantity', row_data.get('total_quantity', 1))),
                    }

                    existing_book = Book.objects.filter(isbn=book_data['isbn']).first()
                    if existing_book:
                        serializer = BookSerializer(existing_book, data=book_data, partial=True)
                        updated_count += 1
                    else:
                        serializer = BookSerializer(data=book_data)
                        created_count += 1

                    if serializer.is_valid():
                        serializer.save()
                    else:
                        errors.append(f'第{row_idx}行: {str(serializer.errors)}')

                except Exception as e:
                    errors.append(f'第{row_idx}行: {str(e)}')

            log_operation(
                request, 'import', 'Book', None, None,
                f'批量导入图书: 新增{created_count}本, 更新{updated_count}本, 错误{len(errors)}条'
            )

            return Response({
                'created_count': created_count,
                'updated_count': updated_count,
                'errors': errors
            })

        except Exception as e:
            return Response({'detail': f'文件解析失败: {str(e)}'}, status=status.HTTP_400_BAD_REQUEST)

    def perform_create(self, serializer):
        instance = serializer.save()
        log_operation(
            self.request, 'create', 'Book', instance.id, instance.title,
            f'新增图书: {instance.title}'
        )

    def perform_update(self, serializer):
        instance = serializer.save()
        log_operation(
            self.request, 'update', 'Book', instance.id, instance.title,
            f'更新图书: {instance.title}'
        )

    def perform_destroy(self, instance):
        log_operation(
            self.request, 'delete', 'Book', instance.id, instance.title,
            f'删除图书: {instance.title}'
        )
        instance.delete()


class BorrowRecordViewSet(viewsets.ModelViewSet):
    queryset = BorrowRecord.objects.all()
    serializer_class = BorrowRecordSerializer
    permission_classes = [IsAuthenticated, IsAdminOrOwner]
    filter_backends = [filters.SearchFilter, filters.OrderingFilter]
    search_fields = ['reader__username', 'book__title', 'book__isbn']
    ordering_fields = ['borrow_date', 'due_date', 'return_date', 'created_at']

    def get_serializer_context(self):
        context = super().get_serializer_context()
        context['request'] = self.request
        return context

    def get_queryset(self):
        queryset = BorrowRecord.objects.all()
        if self.request.user.role == 'reader':
            queryset = queryset.filter(reader=self.request.user)

        reader_id = self.request.query_params.get('reader_id')
        if reader_id:
            queryset = queryset.filter(reader_id=reader_id)

        book_id = self.request.query_params.get('book_id')
        if book_id:
            queryset = queryset.filter(book_id=book_id)

        record_status = self.request.query_params.get('status')
        if record_status:
            queryset = queryset.filter(status=record_status)

        return queryset.order_by('-borrow_date')

    @action(detail=False, methods=['post'], permission_classes=[IsAdminRole])
    def borrow(self, request):
        serializer = BorrowBookSerializer(data=request.data, context={'request': request})
        if serializer.is_valid():
            record = serializer.save()

            log_operation(
                request, 'borrow', 'BorrowRecord', record.id, record.book.title,
                f'借书: {record.reader.username} - {record.book.title}'
            )

            return Response(BorrowRecordSerializer(record, context={'request': request}).data, status=status.HTTP_201_CREATED)
        return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)

    @action(detail=False, methods=['post'], permission_classes=[IsAdminRole])
    def return_book(self, request):
        record_id = request.data.get('record_id')
        try:
            record = BorrowRecord.objects.get(id=record_id, status__in=['borrowed', 'overdue'])
        except BorrowRecord.DoesNotExist:
            return Response({'detail': '借阅记录不存在或已归还'}, status=status.HTTP_400_BAD_REQUEST)

        serializer = ReturnBookSerializer(data=request.data)
        if serializer.is_valid():
            record = serializer.update(record, serializer.validated_data)

            log_operation(
                request, 'return', 'BorrowRecord', record.id, record.book.title,
                f'还书: {record.reader.username} - {record.book.title}'
            )

            return Response(BorrowRecordSerializer(record, context={'request': request}).data)
        return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)

    @action(detail=False, methods=['get'], permission_classes=[IsAuthenticated])
    def my_borrowing(self, request):
        records = BorrowRecord.objects.filter(
            reader=request.user,
            status__in=['borrowed', 'overdue']
        ).order_by('-borrow_date')
        page = self.paginate_queryset(records)
        if page is not None:
            serializer = self.get_serializer(page, many=True)
            return self.get_paginated_response(serializer.data)
        serializer = self.get_serializer(records, many=True)
        return Response(serializer.data)

    @action(detail=False, methods=['get'], permission_classes=[IsAuthenticated])
    def my_borrows(self, request):
        records = BorrowRecord.objects.filter(
            reader=request.user
        ).order_by('-borrow_date')
        
        status = request.query_params.get('status')
        if status:
            records = records.filter(status=status)
        
        page = self.paginate_queryset(records)
        if page is not None:
            serializer = self.get_serializer(page, many=True)
            return self.get_paginated_response(serializer.data)
        serializer = self.get_serializer(records, many=True)
        return Response(serializer.data)

    @action(detail=False, methods=['get'], permission_classes=[IsAuthenticated])
    def my_history(self, request):
        records = BorrowRecord.objects.filter(
            reader=request.user,
            status='returned'
        ).order_by('-return_date')
        page = self.paginate_queryset(records)
        if page is not None:
            serializer = self.get_serializer(page, many=True)
            return self.get_paginated_response(serializer.data)
        serializer = self.get_serializer(records, many=True)
        return Response(serializer.data)

    @action(detail=True, methods=['post'], permission_classes=[IsAuthenticated])
    def renew(self, request, pk=None):
        borrow_record = self.get_object()
        serializer = CreateRenewSerializer(
            data={'borrow_record_id': borrow_record.id, 'renew_days': request.data.get('renew_days', 30)},
            context={'request': request}
        )
        if serializer.is_valid():
            renew_record = serializer.save()

            log_operation(
                request, 'renew', 'RenewRecord', renew_record.id, borrow_record.book.title,
                f'申请续借: {borrow_record.book.title}'
            )

            return Response(RenewRecordSerializer(renew_record).data, status=status.HTTP_201_CREATED)
        return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)


class ReservationViewSet(viewsets.ModelViewSet):
    queryset = Reservation.objects.all()
    serializer_class = ReservationSerializer
    permission_classes = [IsAuthenticated, IsAdminOrOwner]
    filter_backends = [filters.SearchFilter, filters.OrderingFilter]
    search_fields = ['reader__username', 'book__title', 'book__isbn']
    ordering_fields = ['reservation_date', 'expire_date', 'queue_position']

    def get_queryset(self):
        queryset = Reservation.objects.all()
        if self.request.user.role == 'reader':
            queryset = queryset.filter(reader=self.request.user)

        reader_id = self.request.query_params.get('reader_id')
        if reader_id:
            queryset = queryset.filter(reader_id=reader_id)

        book_id = self.request.query_params.get('book_id')
        if book_id:
            queryset = queryset.filter(book_id=book_id)

        reservation_status = self.request.query_params.get('status')
        if reservation_status:
            queryset = queryset.filter(status=reservation_status)

        return queryset.order_by('-reservation_date')

    @action(detail=True, methods=['post'], permission_classes=[IsAuthenticated])
    def cancel(self, request, pk=None):
        reservation = self.get_object()

        if request.user.role == 'reader' and reservation.reader != request.user:
            return Response({'detail': '无权限取消他人预约'}, status=status.HTTP_403_FORBIDDEN)

        if reservation.status not in ['pending', 'available']:
            return Response({'detail': '该预约状态不允许取消'}, status=status.HTTP_400_BAD_REQUEST)

        reservation.status = 'cancelled'
        reservation.save()

        book = reservation.book
        book.reservation_count = F('reservation_count') - 1
        book.save()
        book.refresh_from_db()
        book.update_status()

        log_operation(
            request, 'update', 'Reservation', reservation.id, reservation.book.title,
            f'取消预约: {reservation.book.title}'
        )

        return Response(ReservationSerializer(reservation).data)

    @action(detail=False, methods=['get'], permission_classes=[IsAuthenticated])
    def my_reservations(self, request):
        reservations = Reservation.objects.filter(
            reader=request.user,
            status__in=['pending', 'available']
        ).order_by('-reservation_date')
        page = self.paginate_queryset(reservations)
        if page is not None:
            serializer = self.get_serializer(page, many=True)
            return self.get_paginated_response(serializer.data)
        serializer = self.get_serializer(reservations, many=True)
        return Response(serializer.data)


class RenewRecordViewSet(viewsets.ModelViewSet):
    queryset = RenewRecord.objects.all()
    serializer_class = RenewRecordSerializer
    permission_classes = [IsAuthenticated, IsAdminOrOwner]
    filter_backends = [filters.SearchFilter, filters.OrderingFilter]
    search_fields = ['reader__username', 'book__title', 'book__isbn']
    ordering_fields = ['request_date', 'review_date']

    def get_queryset(self):
        queryset = RenewRecord.objects.all()
        if self.request.user.role == 'reader':
            queryset = queryset.filter(reader=self.request.user)

        reader_id = self.request.query_params.get('reader_id')
        if reader_id:
            queryset = queryset.filter(reader_id=reader_id)

        book_id = self.request.query_params.get('book_id')
        if book_id:
            queryset = queryset.filter(book_id=book_id)

        renew_status = self.request.query_params.get('status')
        if renew_status:
            queryset = queryset.filter(status=renew_status)

        return queryset.order_by('-request_date')

    @action(detail=False, methods=['post'], permission_classes=[IsAdminRole])
    def review(self, request):
        renew_id = request.data.get('renew_id')
        approved = request.data.get('approved')
        review_remark = request.data.get('review_remark', '')

        try:
            renew = RenewRecord.objects.get(id=renew_id, status='pending')
        except RenewRecord.DoesNotExist:
            return Response({'detail': '续借记录不存在或已审核'}, status=status.HTTP_400_BAD_REQUEST)

        serializer = ReviewRenewSerializer(
            data={'renew_id': renew_id, 'approved': approved, 'review_remark': review_remark},
            context={'request': request}
        )
        if serializer.is_valid():
            renew = serializer.update(renew, serializer.validated_data)

            log_operation(
                request, 'renew', 'RenewRecord', renew.id, renew.book.title,
                f'{"通过" if approved else "拒绝"}续借: {renew.reader.username} - {renew.book.title}'
            )

            return Response(RenewRecordSerializer(renew).data)
        return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)

    @action(detail=False, methods=['get'], permission_classes=[IsAuthenticated])
    def my_renews(self, request):
        renews = RenewRecord.objects.filter(
            reader=request.user
        ).order_by('-request_date')
        page = self.paginate_queryset(renews)
        if page is not None:
            serializer = self.get_serializer(page, many=True)
            return self.get_paginated_response(serializer.data)
        serializer = self.get_serializer(renews, many=True)
        return Response(serializer.data)


class FineRecordViewSet(viewsets.ModelViewSet):
    queryset = FineRecord.objects.all()
    serializer_class = FineRecordSerializer
    permission_classes = [IsAuthenticated, IsAdminOrOwner]
    filter_backends = [filters.SearchFilter, filters.OrderingFilter]
    search_fields = ['reader__username', 'book__title', 'book__isbn']
    ordering_fields = ['created_at', 'paid_date', 'amount']

    def get_queryset(self):
        queryset = FineRecord.objects.all()
        if self.request.user.role == 'reader':
            queryset = queryset.filter(reader=self.request.user)

        reader_id = self.request.query_params.get('reader_id')
        if reader_id:
            queryset = queryset.filter(reader_id=reader_id)

        book_id = self.request.query_params.get('book_id')
        if book_id:
            queryset = queryset.filter(book_id=book_id)

        fine_status = self.request.query_params.get('status')
        if fine_status:
            queryset = queryset.filter(status=fine_status)

        return queryset.order_by('-created_at')

    @action(detail=False, methods=['post'], permission_classes=[IsAdminRole])
    def pay(self, request):
        fine_id = request.data.get('fine_id')
        payment_method = request.data.get('payment_method')
        remarks = request.data.get('remarks', '')

        try:
            fine = FineRecord.objects.get(id=fine_id, status='unpaid')
        except FineRecord.DoesNotExist:
            return Response({'detail': '罚款记录不存在或已处理'}, status=status.HTTP_400_BAD_REQUEST)

        serializer = PayFineSerializer(
            data={'fine_id': fine_id, 'payment_method': payment_method, 'remarks': remarks},
            context={'request': request}
        )
        if serializer.is_valid():
            fine = serializer.update(fine, serializer.validated_data)

            log_operation(
                request, 'fine', 'FineRecord', fine.id, fine.book.title,
                f'缴纳罚款: {fine.reader.username} - ¥{fine.amount}'
            )

            return Response(FineRecordSerializer(fine).data)
        return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)

    @action(detail=True, methods=['post'], permission_classes=[IsSuperAdmin])
    def waive(self, request, pk=None):
        fine = self.get_object()
        if fine.status != 'unpaid':
            return Response({'detail': '该罚款已处理'}, status=status.HTTP_400_BAD_REQUEST)

        waiver_reason = request.data.get('waiver_reason', '')

        fine.status = 'waived'
        fine.waiver_reason = waiver_reason
        fine.paid_by = request.user
        fine.paid_date = timezone.now()
        fine.save()

        borrow_record = fine.borrow_record
        borrow_record.fine_paid = True
        borrow_record.fine_paid_at = timezone.now()
        borrow_record.save()

        log_operation(
            request, 'fine', 'FineRecord', fine.id, fine.book.title,
            f'减免罚款: {fine.reader.username} - ¥{fine.amount}, 原因: {waiver_reason}'
        )

        return Response(FineRecordSerializer(fine).data)

    @action(detail=False, methods=['get'], permission_classes=[IsAuthenticated])
    def my_fines(self, request):
        fines = FineRecord.objects.filter(
            reader=request.user
        ).order_by('-created_at')
        page = self.paginate_queryset(fines)
        if page is not None:
            serializer = self.get_serializer(page, many=True)
            return self.get_paginated_response(serializer.data)
        serializer = self.get_serializer(fines, many=True)
        return Response(serializer.data)

    @action(detail=False, methods=['get'], permission_classes=[IsAuthenticated])
    def my_unpaid_fines(self, request):
        fines = FineRecord.objects.filter(
            reader=request.user,
            status='unpaid'
        ).order_by('-created_at')
        total_amount = fines.aggregate(total=Sum('amount'))['total'] or Decimal('0')
        return Response({
            'fines': FineRecordSerializer(fines, many=True).data,
            'total_amount': str(total_amount),
            'count': fines.count()
        })


class OperationLogViewSet(viewsets.ReadOnlyModelViewSet):
    queryset = OperationLog.objects.all()
    serializer_class = OperationLogSerializer
    permission_classes = [IsAdminRole]
    filter_backends = [filters.SearchFilter, filters.OrderingFilter]
    search_fields = ['operator__username', 'operation_type', 'target_model', 'target_name', 'description']
    ordering_fields = ['created_at']

    def get_queryset(self):
        queryset = OperationLog.objects.all()

        operator_id = self.request.query_params.get('operator_id')
        if operator_id:
            queryset = queryset.filter(operator_id=operator_id)

        operation_type = self.request.query_params.get('operation_type')
        if operation_type:
            queryset = queryset.filter(operation_type=operation_type)

        target_model = self.request.query_params.get('target_model')
        if target_model:
            queryset = queryset.filter(target_model=target_model)

        start_date = self.request.query_params.get('start_date')
        if start_date:
            queryset = queryset.filter(created_at__date__gte=start_date)

        end_date = self.request.query_params.get('end_date')
        if end_date:
            queryset = queryset.filter(created_at__date__lte=end_date)

        return queryset.order_by('-created_at')


class SystemConfigViewSet(viewsets.ModelViewSet):
    queryset = SystemConfig.objects.all()
    serializer_class = SystemConfigSerializer
    permission_classes = [IsSuperAdmin]

    def get_queryset(self):
        return SystemConfig.objects.all().order_by('key')

    def perform_update(self, serializer):
        instance = serializer.save()
        log_operation(
            self.request, 'update', 'SystemConfig', instance.id, instance.key,
            f'更新系统配置: {instance.key} = {instance.value}'
        )


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def dashboard_stats(request):
    user = request.user
    is_admin = user.role in ['super_admin', 'admin']

    if is_admin:
        total_books = Book.objects.count()
        available_books = Book.objects.filter(available_quantity__gt=0).count()
        total_readers = User.objects.filter(role='reader').count()
        borrowed_count = BorrowRecord.objects.filter(status__in=['borrowed', 'overdue']).count()
        overdue_count = BorrowRecord.objects.filter(status='overdue').count()
        pending_renewal_count = RenewRecord.objects.filter(status='pending').count()
        unpaid_fine = FineRecord.objects.filter(status='unpaid').aggregate(total=Sum('amount'))['total'] or 0
        today = timezone.now().date()
        today_borrowed = BorrowRecord.objects.filter(borrow_date__date=today).count()

        return Response({
            'total_books': total_books,
            'available_books': available_books,
            'total_readers': total_readers,
            'borrowed_count': borrowed_count,
            'overdue_count': overdue_count,
            'pending_renewal_count': pending_renewal_count,
            'unpaid_fine': float(unpaid_fine),
            'today_borrowed': today_borrowed,
        })
    else:
        total_borrowed = BorrowRecord.objects.filter(reader=user).count()
        current_borrowed = BorrowRecord.objects.filter(reader=user, status__in=['borrowed', 'overdue', 'renewed']).count()
        returned_count = BorrowRecord.objects.filter(reader=user, status='returned').count()
        overdue_count = BorrowRecord.objects.filter(reader=user, status='overdue').count()
        unpaid_fine = FineRecord.objects.filter(reader=user, status='unpaid').aggregate(total=Sum('amount'))['total'] or 0
        active_reservations = Reservation.objects.filter(reader=user, status__in=['pending', 'available']).count()

        return Response({
            'total_borrowed': total_borrowed,
            'current_borrowed': current_borrowed,
            'returned_count': returned_count,
            'overdue_count': overdue_count,
            'unpaid_fine': float(unpaid_fine),
            'active_reservations': active_reservations,
        })


@api_view(['GET'])
@permission_classes([IsAdminRole])
def borrow_stats(request):
    days = int(request.query_params.get('days', 30))
    now = timezone.now()
    start_date = now - timedelta(days=days)

    borrow_records = BorrowRecord.objects.filter(
        borrow_date__gte=start_date
    ).values_list('borrow_date', 'return_date', 'book__title', 'book__isbn', 
                   'book__category__name', 'reader__username', 'reader__first_name', 
                   'reader__last_name', 'status', 'due_date')

    date_list = [(start_date + timedelta(days=i)).date() for i in range(days + 1)]
    borrow_trend_dict = {d: 0 for d in date_list}
    return_trend_dict = {d: 0 for d in date_list}
    overdue_trend_dict = {d: 0 for d in date_list}

    borrow_category_dict = {}
    book_count_dict = {}
    reader_count_dict = {}

    for rec in borrow_records:
        borrow_date, return_date, book_title, book_isbn, category, reader_username, \
        reader_first, reader_last, status, due_date = rec
        
        if borrow_date:
            b_date = borrow_date.date()
            if b_date in borrow_trend_dict:
                borrow_trend_dict[b_date] += 1
        
        if return_date:
            r_date = return_date.date()
            if r_date in return_trend_dict:
                return_trend_dict[r_date] += 1
        
        if due_date and status in ['overdue', 'returned']:
            d_date = due_date.date()
            if d_date in overdue_trend_dict:
                overdue_trend_dict[d_date] += 1
        
        if category:
            borrow_category_dict[category] = borrow_category_dict.get(category, 0) + 1
        
        book_key = (book_title, book_isbn)
        book_count_dict[book_key] = book_count_dict.get(book_key, 0) + 1
        
        reader_key = (reader_username, reader_first, reader_last)
        reader_count_dict[reader_key] = reader_count_dict.get(reader_key, 0) + 1

    borrow_trend = [
        {'date': d.isoformat(), 'borrow_count': borrow_trend_dict[d], 'return_count': return_trend_dict[d]}
        for d in sorted(borrow_trend_dict.keys())
    ]

    return_trend = [
        {'date': d.isoformat(), 'count': return_trend_dict[d]}
        for d in sorted(return_trend_dict.keys())
    ]

    overdue_trend = [
        {'date': d.isoformat(), 'overdue_count': overdue_trend_dict[d], 'new_overdue': 0}
        for d in sorted(overdue_trend_dict.keys())
    ]

    book_categories = Book.objects.values_list('category__name').annotate(
        count=Count('id')
    ).order_by('-count')
    
    category_distribution = [
        {'name': name, 'count': count} for name, count in book_categories
    ]

    borrow_category_distribution = [
        {'name': k, 'count': v} 
        for k, v in sorted(borrow_category_dict.items(), key=lambda x: -x[1])
    ]

    top_books = sorted(book_count_dict.items(), key=lambda x: -x[1])[:10]
    top_books_list = [
        {'title': title, 'isbn': isbn, 'borrow_count': count}
        for (title, isbn), count in top_books
    ]

    top_readers = sorted(reader_count_dict.items(), key=lambda x: -x[1])[:10]
    top_readers_list = [
        {'reader_name': f'{first}{last}', 'username': username, 'borrow_count': count}
        for (username, first, last), count in top_readers
    ]

    overdue_records = BorrowRecord.objects.filter(
        status='overdue'
    ).values_list('due_date', flat=True)
    
    overdue_by_days = {}
    for due_date in overdue_records:
        if due_date:
            diff_days = (now.date() - due_date).days
            overdue_by_days[diff_days] = overdue_by_days.get(diff_days, 0) + 1
    
    current_overdue_by_days = [
        {'overdue_days': k, 'count': v} 
        for k, v in sorted(overdue_by_days.items())
    ]

    return Response({
        'borrow_trend': borrow_trend,
        'return_trend': return_trend,
        'monthly_borrow': [],
        'book_categories': category_distribution,
        'borrow_categories': borrow_category_distribution,
        'hot_books': top_books_list,
        'reader_rank': top_readers_list,
        'overdue_trend': overdue_trend,
        'current_overdue_by_days': current_overdue_by_days,
    })


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def reader_stats(request):
    reader = request.user

    total_borrowed = BorrowRecord.objects.filter(reader=reader).count()
    current_borrowed = BorrowRecord.objects.filter(reader=reader, status__in=['borrowed', 'overdue']).count()
    total_overdue = BorrowRecord.objects.filter(reader=reader, status='overdue').count()
    total_returned = BorrowRecord.objects.filter(reader=reader, status='returned').count()

    total_fines = FineRecord.objects.filter(reader=reader).aggregate(total=Sum('amount'))['total'] or 0
    unpaid_fines = FineRecord.objects.filter(reader=reader, status='unpaid').aggregate(total=Sum('amount'))['total'] or 0

    active_reservations = Reservation.objects.filter(reader=reader, status__in=['pending', 'available']).count()
    pending_renews = RenewRecord.objects.filter(reader=reader, status='pending').count()

    days = int(request.query_params.get('days', 90))
    start_date = timezone.now() - timedelta(days=days)

    borrow_history = BorrowRecord.objects.filter(
        reader=reader,
        borrow_date__gte=start_date
    ).annotate(
        date=TruncDate('borrow_date')
    ).values('date').annotate(
        count=Count('id')
    ).order_by('date')

    category_preference = BorrowRecord.objects.filter(
        reader=reader,
        borrow_date__gte=start_date
    ).values(
        'book__category__name'
    ).annotate(
        count=Count('id')
    ).order_by('-count')

    return Response({
        'total_borrowed': total_borrowed,
        'current_borrowed': current_borrowed,
        'total_overdue': total_overdue,
        'total_returned': total_returned,
        'total_fines': str(total_fines),
        'unpaid_fines': str(unpaid_fines),
        'active_reservations': active_reservations,
        'pending_renews': pending_renews,
        'borrow_history': list(borrow_history),
        'category_preference': list(category_preference),
    })
